"""Word (.docx) Preprocessor.
Extracts document properties, headings, body paragraphs, and structured tables using python-docx.
Normalizes whitespace and chunks by sections and tables.
"""

import logging
import time
from pathlib import Path
from typing import Any, Dict, List
import docx

from app.services.preprocessing.base import BasePreprocessor, PreprocessingChunk, PreprocessingResult

logger = logging.getLogger("securedoc.preprocessing.word")


class WordPreprocessor(BasePreprocessor):
    """Preprocessor for Microsoft Word (.docx) documents."""

    def process(self, doc_id: str, file_path: str, filename: str) -> PreprocessingResult:
        start_time = time.time()
        try:
            doc = docx.Document(file_path)
            core_props = doc.core_properties

            sections_data: List[str] = []
            chunks: List[Dict[str, Any]] = []
            chunk_idx = 1
            before_paras_count = len(doc.paragraphs)
            before_tables_count = len(doc.tables)
            raw_chars_before = 0

            # 1. Process Paragraphs & Headings
            current_section_title = "Document Body"
            current_section_lines: List[str] = []

            for p in doc.paragraphs:
                p_text = p.text or ""
                raw_chars_before += len(p_text)
                cleaned_p = p_text.strip()
                if not cleaned_p:
                    continue

                # Detect Headings
                is_heading = p.style and ("heading" in p.style.name.lower() or "title" in p.style.name.lower())
                if is_heading:
                    if current_section_lines:
                        sec_text = "\n".join(current_section_lines)
                        sections_data.append(sec_text)
                        chunks.append(
                            PreprocessingChunk(
                                chunk_id=chunk_idx,
                                text=sec_text,
                                location=f"section: {current_section_title}",
                                char_count=len(sec_text),
                            ).to_dict()
                        )
                        chunk_idx += 1
                        current_section_lines = []
                    current_section_title = cleaned_p
                    current_section_lines.append(f"## {cleaned_p}")
                else:
                    current_section_lines.append(cleaned_p)

            if current_section_lines:
                sec_text = "\n".join(current_section_lines)
                sections_data.append(sec_text)
                chunks.append(
                    PreprocessingChunk(
                        chunk_id=chunk_idx,
                        text=sec_text,
                        location=f"section: {current_section_title}",
                        char_count=len(sec_text),
                    ).to_dict()
                )
                chunk_idx += 1

            # 2. Process Tables
            table_idx = 1
            for table in doc.tables:
                table_rows: List[str] = []
                for row in table.rows:
                    row_cells = [cell.text.strip().replace("\n", " ") for cell in row.cells]
                    # deduplicate adjacent merged cells in docx
                    deduped_cells = []
                    for c in row_cells:
                        if not deduped_cells or c != deduped_cells[-1]:
                            deduped_cells.append(c)
                    table_rows.append(" | ".join(deduped_cells))
                    raw_chars_before += sum(len(c) for c in row_cells)

                if table_rows:
                    table_md = "\n".join(table_rows)
                    sections_data.append(f"[Table {table_idx}]\n{table_md}")
                    chunks.append(
                        PreprocessingChunk(
                            chunk_id=chunk_idx,
                            text=table_md,
                            location=f"table: {table_idx}",
                            char_count=len(table_md),
                        ).to_dict()
                    )
                    chunk_idx += 1
                    table_idx += 1

            cleaned_content = "\n\n".join(sections_data).strip()
            raw_chars_after = len(cleaned_content)
            duration_ms = (time.time() - start_time) * 1000

            return PreprocessingResult(
                doc_id=doc_id,
                data_type="word",
                cleaned_content=cleaned_content,
                chunks=chunks,
                metadata={
                    "filename": filename,
                    "title": getattr(core_props, "title", "") or "",
                    "author": getattr(core_props, "author", "") or "",
                    "paragraphs_count": before_paras_count,
                    "tables_count": before_tables_count,
                    "created": str(getattr(core_props, "created", "") or ""),
                },
                stats={
                    "before_chars": raw_chars_before,
                    "after_chars": raw_chars_after,
                    "paragraphs_count": before_paras_count,
                    "tables_count": before_tables_count,
                    "chunks_count": len(chunks),
                },
                status="success",
                processing_time_ms=round(duration_ms, 2),
            )

        except Exception as e:
            logger.error("Word preprocessing failed for %s: %s", doc_id, e, exc_info=True)
            duration_ms = (time.time() - start_time) * 1000
            return PreprocessingResult(
                doc_id=doc_id,
                data_type="word",
                cleaned_content="",
                status="failed",
                error_message=f"Word preprocessing error: {str(e)}",
                processing_time_ms=round(duration_ms, 2),
            )
